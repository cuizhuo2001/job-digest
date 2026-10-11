#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
把求职信 .md 转成 .docx。

用法：
    python md2docx.py GoldmanSachs_QuantStrats_NewAnalyst.md
    python md2docx.py xxx.md -o 自定义输出名.docx

规则：
- 取第一个 `---` 和第二个 `---` 之间的正文（前后是元信息和设计说明，不进 docx）
- 空行分段；**粗体** → 加粗；`反引号` → 去掉反引号
- 排版对齐港版简历：Arial、10.5pt、姓名用深蓝粗体
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

NAVY = RGBColor(0x1A, 0x3C, 0x6E)
FONT = "Arial"
SIZE = Pt(10.5)


def extract_body(md_text: str) -> list[str]:
    """取两个 --- 之间的正文段落。"""
    parts = re.split(r"(?m)^---\s*$", md_text)
    if len(parts) < 3:
        raise SystemExit("找不到 `---` 分隔线，无法定位正文。请检查 md 结构。")
    body = parts[1]

    paragraphs = []
    for block in re.split(r"\n\s*\n", body.strip()):
        block = block.strip()
        if not block:
            continue
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        # 含整行粗体的块（如签名「**你的名字**」+ 联系方式）保留换行，其余连排
        if len(lines) > 1 and any(re.fullmatch(r"\*\*.+\*\*", l) for l in lines):
            paragraphs.append("\n".join(lines))
        else:
            paragraphs.append(" ".join(lines))
    return paragraphs


def add_paragraph(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.25

    for i, line in enumerate(text.split("\n")):
        if i:
            p.add_run().add_break()
        # **粗体** 与 `代码` 混排
        for token in re.split(r"(\*\*.+?\*\*|`.+?`)", line):
            if not token:
                continue
            bold = token.startswith("**") and token.endswith("**")
            clean = token[2:-2] if bold else (token[1:-1] if token.startswith("`") else token)

            run = p.add_run(clean)
            run.font.name = FONT
            run.font.size = SIZE
            if bold:
                run.bold = True
                run.font.color.rgb = NAVY
    return p


def main():
    args = [a for a in sys.argv[1:] if a != "-o"]
    if not args:
        raise SystemExit(__doc__)

    src = Path(args[0])
    if not src.exists():
        src = Path(__file__).parent / args[0]
    if not src.exists():
        raise SystemExit(f"找不到文件：{args[0]}")

    out = None
    if "-o" in sys.argv:
        i = sys.argv.index("-o")
        out = Path(sys.argv[i + 1])
    out = out or src.with_suffix(".docx")

    md = src.read_text(encoding="utf-8")
    paragraphs = extract_body(md)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = SIZE

    for text in paragraphs:
        add_paragraph(doc, text)

    # 让正文左对齐、页边距正常
    for section in doc.sections:
        section.left_margin = Pt(72)
        section.right_margin = Pt(72)
        section.top_margin = Pt(60)
        section.bottom_margin = Pt(60)

    doc.save(out)
    print(f"OK -> {out}  ({len(paragraphs)} 段)")


if __name__ == "__main__":
    main()
